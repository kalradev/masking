import re
import shutil
import subprocess
import uuid
import zipfile
from functools import lru_cache
from pathlib import Path

import pymupdf as fitz
from flask import Flask, jsonify, render_template, request, send_file
from pdf2docx import Converter
from werkzeug.utils import secure_filename

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 150 * 1024 * 1024  # 150 MB batch
MAX_FILES = 25

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "storage" / "uploads"
OUTPUT_DIR = BASE_DIR / "storage" / "masked"
PAGES_DIR = BASE_DIR / "storage" / "pages"
CONVERT_UPLOAD_DIR = BASE_DIR / "storage" / "convert_uploads"
CONVERT_OUTPUT_DIR = BASE_DIR / "storage" / "converted"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
PAGES_DIR.mkdir(parents=True, exist_ok=True)
CONVERT_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
CONVERT_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

JOB_ID_RE = re.compile(r"^[a-f0-9]{32}$")
PAGE_ZOOM = 2.0

PII_PATTERNS = {
    "email": re.compile(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b"),
    "phone": re.compile(
        r"(?:\+?\d{1,3}[-.\s]?)?(?:\(?\d{2,4}\)?[-.\s]?)?\d{3,4}[-.\s]?\d{3,4}\b"
    ),
    "ssn": re.compile(r"\b\d{3}-\d{2}-\d{4}\b"),
    "credit_card": re.compile(r"\b(?:\d[ -]*?){13,19}\b"),
    "aadhaar": re.compile(r"\b\d{4}\s?\d{4}\s?\d{4}\b"),
    "pan": re.compile(r"\b[A-Z]{5}\d{4}[A-Z]\b"),
    "ip": re.compile(r"\b(?:\d{1,3}\.){3}\d{1,3}\b"),
    "ifsc": re.compile(r"\b[A-Z]{4}0[A-Z0-9]{6}\b"),
    "passport": re.compile(r"\b[A-PR-WY][1-9]\d{6}\b"),
    "gstin": re.compile(r"\b\d{2}[A-Z]{5}\d{4}[A-Z][A-Z0-9]Z[A-Z0-9]\b"),
}

DEFAULT_TYPES = ["email", "phone", "ssn", "credit_card", "aadhaar", "pan", "ip"]
AUTO_TYPES = DEFAULT_TYPES + ["ifsc", "passport"]

NUMBER_WORDS = (
    r"zero|one|two|three|four|five|six|seven|eight|nine|ten|"
    r"eleven|twelve|thirteen|fourteen|fifteen|sixteen|seventeen|eighteen|nineteen|"
    r"twenty|thirty|forty|fifty|sixty|seventy|eighty|ninety|"
    r"hundred|thousand|lakh|lac|lakhs|crore|crores|million|billion|paise|rupees?"
)

AMOUNT_NUMERIC = re.compile(
    r"(?ix)"
    r"(?:(?:₹|rs\.?|inr|usd|eur|gbp|\$|€|£)\s*)"
    r"(?:"
    r"\d{1,2}(?:,\d{2})+,\d{3}(?:\.\d{1,2})?"
    r"|\d{1,3}(?:,\d{3})+(?:\.\d{1,2})?"
    r"|\d{4,12}(?:\.\d{1,2})?"
    r"|\d+(?:\.\d{1,2})?/-"
    r")"
    r"|"
    r"(?<![\d.])"
    r"(?:"
    r"\d{1,2}(?:,\d{2})+,\d{3}(?:\.\d{1,2})?"
    r"|\d{1,3}(?:,\d{3})+(?:\.\d{1,2})?"
    r"|\d{5,12}(?:\.\d{1,2})?"
    r"|\d{4,12}\.\d{2}"
    r"|\d{3,12}(?:/-)"
    r")"
    r"(?![\d.])"
)

QTY_NEAR = re.compile(
    r"(?i)\b(?:qty|quantity|quantities|total\s+qty|nos\.?|n\.o\.s\.?|pcs\.?|pieces?|units?)\b"
)

AMOUNT_WORDS = re.compile(
    rf"(?ix)"
    rf"(?:(?:rs\.?|rupees?|inr|usd|dollars?|eur|euros?)\s+)"
    rf"(?:(?:{NUMBER_WORDS}|and)[\s\-]+){{0,12}}"
    rf"(?:{NUMBER_WORDS})"
    rf"(?:[\s\-]+(?:{NUMBER_WORDS}|and))*"
    rf"(?:\s+only)?"
    rf"|"
    rf"(?:(?:{NUMBER_WORDS}|and)[\s\-]+){{1,12}}"
    rf"(?:thousand|lakh|lac|lakhs|crore|crores|million|billion|hundred)"
    rf"(?:[\s\-]+(?:{NUMBER_WORDS}|and))*"
    rf"(?:\s+only)?"
)

GST_MONEY = re.compile(
    r"(?ix)"
    r"\b(?:igst|cgst|sgst|ugst|gst)\b"
    r"(?:\s*\(\s*\d+(?:\.\d+)?\s*%\s*\))?"
    r"(?:\s+(?:value|amount|on(?:\s+po)?))?"
    r"(?:\s+(?:for|of|on))?"
    r"(?:\s+po)?"
    r"\s*[:\-]?\s*"
    r"(?:₹|rs\.?|inr)?\s*"
    r"(\d{1,3}(?:,\d{2,3})+(?:\.\d{1,2})?|\d{4,12}(?:\.\d{1,2})?)"
)

AMOUNT_LABEL = re.compile(
    r"(?i)\b("
    r"amount(?:\s+in\s+words)?|total\s+po\s+value|po\s+value|"
    r"total(?:\s+amount)?|grand\s+total|net\s+amount|"
    r"invoice\s+amount|taxable(?:\s+value|\s+amount)?|sub[\s-]?total|"
    r"line\s+total|unit\s+price|price|rate|cost|payable|balance|charges?|discount|"
    r"cgst|sgst|igst|gst\s+amount|gst\s+value|igst\s+value|cgst\s+value|sgst\s+value|"
    r"salary|income"
    r")\s*[:\-]\s*([^\n;|]+)"
)

QTY_LABELED = re.compile(
    r"(?i)\b(?:qty|quantity)\b\s*[:\-]?\s*(\d{1,6}(?:\.\d{1,3})?)"
)

PIN_NEAR = re.compile(r"(?i)\b(?:pin|pincode|postal\s+code)\b")

TABLE_NUMBER = re.compile(r"^[\d,]+(?:\.\d{1,3})?$")


def unique_terms(items: list[str]) -> list[str]:
    unique = []
    seen = set()
    for item in items:
        candidate = " ".join(item.split()).strip(" :-")
        if len(candidate) < 2:
            continue
        key = candidate.lower()
        if key not in seen:
            seen.add(key)
            unique.append(candidate)
    return unique


def mask_keep_last_digits(value: str, keep: int = 4) -> str:
    digits = re.sub(r"\D", "", value)
    if len(digits) <= keep:
        return "*" * max(1, len(value))
    keep_from = len(digits) - keep
    seen = 0
    out = []
    for char in value:
        if char.isdigit():
            out.append(char if seen >= keep_from else "*")
            seen += 1
        else:
            out.append(char)
    return "".join(out)


def partial_mask(value: str, pii_type: str | None, auto: bool) -> str:
    text = value.strip()
    if not auto:
        return "*" * max(1, len(text))

    if pii_type == "amount":
        return "*" * max(3, min(len(text), 18))

    if pii_type == "email" or "@" in text:
        _, _, domain = text.partition("@")
        if domain:
            return f"***@{domain}"

    if pii_type == "phone":
        return mask_keep_last_digits(text, 4)
    if pii_type == "aadhaar":
        return mask_keep_last_digits(text, 4)
    if pii_type == "credit_card":
        return mask_keep_last_digits(text, 4)
    if pii_type == "ssn" and len(text) >= 4:
        return "***-**-" + text[-4:]
    if pii_type == "pan" and len(text) == 10:
        return "*" * 5 + text[5:]
    if pii_type == "ip":
        parts = text.split(".")
        if len(parts) == 4:
            return "***.***.***." + parts[3]
    if pii_type == "ifsc" and len(text) > 4:
        return text[:4] + "*" * (len(text) - 4)
    if pii_type in ("passport", "gstin") and len(text) > 4:
        return "*" * (len(text) - 4) + text[-4:]
    if len(text) <= 2:
        return "*" * len(text)
    return text[0] + "***"


GSTIN_VALUE = re.compile(r"\d{2}[A-Z]{5}\d{4}[A-Z][A-Z0-9]Z[A-Z0-9]", re.I)


def looks_like_gstin(value: str) -> bool:
    compact = re.sub(r"[\s-]", "", value).upper()
    return bool(GSTIN_VALUE.search(compact))


CURRENCY_PREFIX = re.compile(
    r"^(?:₹|rs\.?|rupees?|inr|usd|dollars?|eur|euros?|gbp|\$|€|£)\s*",
    re.I,
)
CURRENCY_ONLY = re.compile(
    r"^(?:₹|rs\.?|rupees?|inr|usd|dollars?|eur|euros?|gbp|\$|€|£)$",
    re.I,
)


def amount_text_to_redact(value: str, allow_quantity: bool = False) -> str | None:
    text = " ".join(value.split()).strip(" :-()")
    text = CURRENCY_PREFIX.sub("", text).strip(" :-()")
    if not text or CURRENCY_ONLY.fullmatch(text) or looks_like_gstin(text):
        return None
    if not allow_quantity and is_likely_quantity(text):
        return None
    if not re.search(r"\d", text) and not AMOUNT_WORDS.search(text):
        return None
    if not allow_quantity and not is_amount_candidate(text) and not AMOUNT_WORDS.search(text):
        return None
    if allow_quantity and looks_like_gstin(text):
        return None
    return text


def is_likely_quantity(value: str) -> bool:
    text = value.strip()
    if QTY_NEAR.search(text):
        return True
    return bool(re.fullmatch(r"\d{1,3}(?:\.\d{1,2})?", text))


def near_quantity_label(text: str, start: int, end: int) -> bool:
    window = text[max(0, start - 80) : min(len(text), end + 30)]
    return bool(QTY_NEAR.search(window))


def near_pin_label(text: str, start: int, end: int) -> bool:
    window = text[max(0, start - 40) : min(len(text), end + 20)]
    return bool(PIN_NEAR.search(window))


def is_amount_candidate(value: str) -> bool:
    compact = value.strip()
    if len(compact) < 3:
        return False
    if re.fullmatch(r"(?:19|20)\d{2}", compact):
        return False
    if looks_like_gstin(compact):
        return False
    if is_likely_quantity(compact):
        return False
    digits = re.sub(r"\D", "", compact)
    if len(digits) >= 10 and not re.search(r"[,\$\u20b9]|/-|\.\d{2}|rs\.?|inr", compact, re.I):
        return False
    if len(digits) == 8 and not re.search(r"[,]|/-|\.\d", compact):
        return False
    return True


def find_matches(
    text: str, selected_types: list[str], extra_terms: list[str], auto: bool
) -> list[tuple[str, str]]:
    found: list[tuple[str, str]] = []

    if auto:
        for pattern in (AMOUNT_NUMERIC, AMOUNT_WORDS, GST_MONEY):
            for match in pattern.finditer(text):
                if near_pin_label(text, match.start(), match.end()):
                    continue
                raw = match.group(1) if match.lastindex else match.group(0)
                candidate = amount_text_to_redact(raw)
                if not candidate:
                    continue
                found.append((candidate, partial_mask(candidate, "amount", auto)))
        for _label, value in AMOUNT_LABEL.findall(text):
            candidate = amount_text_to_redact(value)
            if not candidate:
                continue
            found.append((candidate, partial_mask(candidate, "amount", auto)))
        for match in QTY_LABELED.finditer(text):
            phrase = match.group(0).strip()
            if phrase:
                found.append((phrase, partial_mask(match.group(1), "amount", auto)))
    else:
        types = selected_types or DEFAULT_TYPES
        for pii_type in types:
            pattern = PII_PATTERNS.get(pii_type)
            if not pattern:
                continue
            for match in pattern.finditer(text):
                candidate = match.group(0).strip()
                if looks_like_gstin(candidate):
                    continue
                if pii_type == "phone" and len(re.sub(r"\D", "", candidate)) < 10:
                    continue
                if pii_type == "credit_card" and len(re.sub(r"\D", "", candidate)) < 13:
                    continue
                if candidate:
                    found.append((candidate, partial_mask(candidate, pii_type, auto)))

    for term in extra_terms:
        if term and term in text:
            found.append((term, partial_mask(term, "amount" if auto else None, auto)))

    unique = []
    seen = set()
    for original, replacement in found:
        key = original.lower()
        if key not in seen:
            seen.add(key)
            unique.append((original, replacement))
    return unique


def redact_with_stars(page, items: list[tuple], fill_stars: bool = True) -> int:
    if not items:
        return 0
    items = merge_star_hits(items)
    bands = []
    for rect, payload in items:
        band = clip_rect_from_neighbors(page, rect)
        page.add_redact_annot(band, fill=None, cross_out=False)
        bands.append((band, payload))
    page.apply_redactions(images=0, graphics=0)
    if fill_stars:
        for band, payload in bands:
            if payload == "" or payload is False:
                continue
            if isinstance(payload, str):
                write_mask_text(page, band, payload)
            else:
                write_stars(page, band, payload, single_line=True)
    return len(items)


def merge_star_hits(items: list[tuple]) -> list[tuple]:
    """Drop duplicate ***** for the same token (label + column hits).

    Only near-identical rects are merged. Different rows that happen to be
    close (IGST row and Total Line Value row) each keep their own stars.
    Wrapped fragments are grouped by the caller, not here.
    """
    stars = [(fitz.Rect(rect), payload) for rect, payload in items if payload]
    extras = [(fitz.Rect(rect), "") for rect, payload in items if not payload]
    kept = []
    used = set()
    for i, (rect_a, payload_a) in enumerate(stars):
        if i in used:
            continue
        for j, (rect_b, _payload_b) in enumerate(stars):
            if j <= i or j in used:
                continue
            close_x = rect_a.x0 < rect_b.x1 + 4 and rect_a.x1 > rect_b.x0 - 4
            same_line = abs(rect_a.y0 - rect_b.y0) < 3 and abs(rect_a.y1 - rect_b.y1) < 4
            if close_x and same_line:
                extras.append((rect_b, ""))
                used.add(j)
        kept.append((rect_a, payload_a))
        used.add(i)
    return kept + extras


def tight_redact_rect(rect: fitz.Rect) -> fitz.Rect:
    """Cover the first line of the cell, where the digits actually sit."""
    height = max(float(rect.height), 1.0)
    line_h = 10.5 if height > 13 else max(6.5, height * 0.78)
    y0 = rect.y0 + min(0.7, height * 0.05)
    y1 = min(rect.y1 - 0.35, y0 + line_h)
    if y1 - y0 < 4:
        y1 = min(rect.y1, y0 + 6)
    return fitz.Rect(rect.x0 + 0.15, y0, rect.x1 - 0.1, y1)


def clip_rect_from_neighbors(page, rect: fitz.Rect) -> fitz.Rect:
    """Shrink the redact box so it only touches this line's glyphs.

    Word boxes of adjacent lines overlap by a couple of points, and MuPDF
    removes any character whose box intersects a redaction. So the band
    must start below the previous line's box and end above the next one.
    """
    center_y = (rect.y0 + rect.y1) / 2
    y0 = rect.y0 + 0.4
    y1 = rect.y1 - 0.4
    for word in page.get_text("words"):
        neighbor = fitz.Rect(word[:4])
        if neighbor.is_empty or neighbor.x1 <= rect.x0 or neighbor.x0 >= rect.x1:
            continue
        neighbor_center = (neighbor.y0 + neighbor.y1) / 2
        if abs(neighbor_center - center_y) <= 3.0:
            continue  # same line (includes the target words themselves)
        if neighbor_center < center_y:
            if neighbor.y1 > y0:
                y0 = max(y0, neighbor.y1 + 0.25)
        else:
            if neighbor.y0 < y1:
                y1 = min(y1, neighbor.y0 - 0.25)
    if y1 - y0 < 1.2:
        y0 = center_y - 0.6
        y1 = center_y + 0.6
    return fitz.Rect(rect.x0 + 0.15, y0, rect.x1 - 0.1, y1)


def star_baseline(rect: fitz.Rect, fontsize: float) -> float:
    """Place stars on the redacted text line."""
    height = max(float(rect.height), 4.0)
    return min(rect.y1 - 0.4, rect.y0 + min(height * 0.78, fontsize * 0.9))


def write_mask_text(page, rect, text: str) -> None:
    height = max(float(rect.height), 4.0)
    width = max(float(rect.width), 4.0)
    fontsize = 8.5
    while fontsize > 6 and fontsize * 0.5 * max(len(text), 1) > width:
        fontsize -= 0.3
    page.insert_text(
        (rect.x0 + 0.6, star_baseline(rect, fontsize)),
        text,
        fontname="helv",
        fontsize=fontsize,
        color=(0, 0, 0),
    )


def write_stars(page, rect, char_count: int | None = None, single_line: bool = False) -> None:
    height = max(float(rect.height), 4.0)
    width = max(float(rect.width), 4.0)
    if single_line or height <= 16:
        fontsize = max(6.0, min(height * 0.55, 9.5))
        cols = max(1, int(width / (fontsize * 0.52)))
        if char_count:
            cols = max(1, min(cols, char_count))
        page.insert_text(
            (rect.x0 + 0.6, star_baseline(rect, fontsize)),
            "*" * cols,
            fontname="helv",
            fontsize=fontsize,
            color=(0, 0, 0),
        )
        return
    fontsize = max(6.0, min(height * 0.72, 11.0))
    step_x = fontsize * 0.52
    step_y = fontsize * 1.05
    cols = max(1, int(width / step_x))
    y = rect.y0 + fontsize * 0.85
    while y <= rect.y1 - fontsize * 0.15:
        page.insert_text(
            (rect.x0 + 1, y),
            "*" * cols,
            fontname="helv",
            fontsize=fontsize,
            color=(0, 0, 0),
        )
        y += step_y


def display_rect(page, region: dict) -> fitz.Rect:
    """Map 0-1 preview fractions onto the same page area used for the PNG."""
    clip = page.rect
    x = max(0.0, min(1.0, float(region["x"])))
    y = max(0.0, min(1.0, float(region["y"])))
    w = max(0.0, min(1.0, float(region["width"])))
    h = max(0.0, min(1.0, float(region["height"])))
    rect = fitz.Rect(
        clip.x0 + x * clip.width,
        clip.y0 + y * clip.height,
        clip.x0 + (x + w) * clip.width,
        clip.y0 + (y + h) * clip.height,
    )
    return rect.normalize()


def coverage_ratio(glyph: fitz.Rect, selection: fitz.Rect) -> float:
    overlap = glyph & selection
    if overlap.is_empty or abs(glyph) <= 0:
        return 0.0
    return abs(overlap) / abs(glyph)


def words_in_selection(page, selection: fitz.Rect) -> list[fitz.Rect]:
    """Only words whose center is inside, or most of the word sits in the box."""
    hits = []
    for word in page.get_text("words"):
        bbox = fitz.Rect(word[:4])
        if bbox.is_empty:
            continue
        center = fitz.Point((bbox.x0 + bbox.x1) / 2, (bbox.y0 + bbox.y1) / 2)
        if selection.contains(center) or coverage_ratio(bbox, selection) >= 0.55:
            hits.append(bbox)
    return hits


def page_word_items(page) -> list[dict]:
    items = []
    for word in page.get_text("words"):
        items.append(
            {
                "x0": word[0],
                "y0": word[1],
                "x1": word[2],
                "y1": word[3],
                "text": word[4],
                "rect": fitz.Rect(word[:4]),
            }
        )
    return items


def group_word_rows(items: list[dict], tol: float = 5.0) -> list[list[dict]]:
    rows: list[list[dict]] = []
    for item in sorted(items, key=lambda i: (i["y0"], i["x0"])):
        if not rows or abs(item["y0"] - rows[-1][0]["y0"]) > tol:
            rows.append([item])
        else:
            rows[-1].append(item)
    for row in rows:
        row.sort(key=lambda i: i["x0"])
    return rows


def header_kind(parts: list[str]) -> str | None:
    phrase = " ".join(parts).lower()
    phrase = re.sub(r"[^a-z0-9\s]+", " ", phrase)
    phrase = " ".join(phrase.split())
    if phrase in {"qty", "quantity", "qty."}:
        return "qty"
    if phrase in {
        "unit price",
        "price",
        "rate",
        "line total",
        "amount",
        "net amount",
        "taxable",
        "taxable value",
        "taxable amount",
        "gst",
        "cgst",
        "sgst",
        "igst",
        "gst amount",
        "total",
    }:
        return "amount"
    return None


def header_columns(row: list[dict]) -> list[tuple[float, float, str]]:
    columns: list[list] = []
    index = 0
    while index < len(row):
        one = [row[index]["text"]]
        two = one + [row[index + 1]["text"]] if index + 1 < len(row) else None
        kind = header_kind(two) if two else None
        span = 2 if kind else 1
        if not kind:
            kind = header_kind(one)
            span = 1
        if kind:
            left = row[index]["x0"]
            right = row[index + span - 1]["x1"]
            columns.append([left, right, kind])
            index += span
        else:
            index += 1
    result = []
    for i, (left, right, kind) in enumerate(columns):
        next_left = columns[i + 1][0] if i + 1 < len(columns) else None
        col_right = ((right + next_left) / 2) if next_left is not None else right + 80
        result.append((left - 8, col_right, kind))
    return result


def find_table_amount_rects(page) -> list[tuple[fitz.Rect, str]]:
    """Mask Qty / Unit Price / Line Total / GST cells using header columns."""
    items = page_word_items(page)
    if not items:
        return []
    hits = []
    for row in group_word_rows(items):
        joined = " ".join(word["text"] for word in row).lower()
        has_qty = bool(re.search(r"\bqty\.?\b|\bquantity\b", joined))
        has_money = bool(re.search(r"\b(?:price|amount|total|rate|gst|cgst|sgst|igst)\b", joined))
        if not (has_qty and has_money):
            continue
        columns = header_columns(row)
        if not columns:
            continue
        header_bottom = max(word["y1"] for word in row)
        for word in items:
            if word["y0"] < header_bottom - 1:
                continue
            token = word["text"].replace(" ", "").replace("₹", "")
            if not TABLE_NUMBER.fullmatch(token):
                continue
            digits = re.sub(r"\D", "", token)
            if not digits or looks_like_gstin(token):
                continue
            if re.fullmatch(r"(?:19|20)\d{2}", digits):
                continue
            center_x = (word["x0"] + word["x1"]) / 2
            for left, right, kind in columns:
                if center_x < left or center_x > right:
                    continue
                stars = partial_mask(token, "amount", True)
                hits.append((word["rect"], stars))
                break
    unique = []
    seen = set()
    for rect, stars in hits:
        key = (round(rect.x0, 1), round(rect.y0, 1), round(rect.x1, 1), round(rect.y1, 1))
        if key not in seen:
            seen.add(key)
            unique.append((rect, stars))
    return unique


MONEY_TOKEN = re.compile(r"^[\d,]+(?:\.\d*)?$")
INR_TOKEN = re.compile(r"^\(?INR\)?\.?$|^₹$", re.I)
GST_LABEL = re.compile(r"(?i)^(cgst|sgst|igst)(_|\d|$)")
AMOUNT_WORD_TOKEN = re.compile(
    rf"(?i)^(?:{NUMBER_WORDS}|and|only|[a-z]+(?:-[a-z]+)+)$"
)

AIRTEL_WORDS_LABELS = (
    ("total", "purchase", "order", "value", "in", "words"),
    ("total", "po", "value", "in", "words"),
)
AIRTEL_MONEY_LABELS = (
    ("total", "purchase", "order", "value"),
    ("total", "po", "value"),
    ("total", "line", "value"),
)
AIRTEL_TERMS_LABELS = (("payment", "terms"),)


def norm_token(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", (text or "").lower())


def is_inr_token(text: str) -> bool:
    return bool(INR_TOKEN.fullmatch((text or "").strip()))


def is_money_token(text: str) -> bool:
    token = (text or "").replace(" ", "").replace("₹", "").strip()
    return bool(MONEY_TOKEN.fullmatch(token)) and not looks_like_gstin(token)


LABEL_TOKENS = {
    "total",
    "purchase",
    "order",
    "value",
    "po",
    "in",
    "words",
    "line",
    "payment",
    "terms",
    "inwords",
}

GROUP_STARS = "*****"


def is_amount_word_token(text: str) -> bool:
    raw = (text or "").strip("():,.")
    token = norm_token(text)
    if not raw or not token or is_inr_token(raw) or is_inr_token(text or ""):
        return False
    if token in LABEL_TOKENS:
        return False
    if AMOUNT_WORD_TOKEN.fullmatch(raw):
        return True
    if re.fullmatch(r"[A-Z][A-Z\-]{2,}", raw) and token not in LABEL_TOKENS:
        return bool(AMOUNT_WORD_TOKEN.fullmatch(raw.lower()) or "-" in raw)
    return False


def skip_label_punct(words: list[dict], index: int) -> int:
    while index < len(words) and words[index]["text"].strip() in {":", "-", "–"}:
        index += 1
    return index


def match_label_at(words: list[dict], start: int, seq: tuple[str, ...]) -> int | None:
    target = "".join(seq)
    acc = ""
    index = start
    while index < len(words) and len(acc) < len(target):
        part = norm_token(words[index]["text"])
        if not part:
            index += 1
            continue
        acc += part
        index += 1
        if not target.startswith(acc) and acc != target:
            return None
    if acc != target:
        return None
    return skip_label_punct(words, index)


def rest_has_in_words(words: list[dict], index: int) -> bool:
    blob = "".join(norm_token(word["text"]) for word in words[index : index + 4])
    return blob.startswith("inwords") or blob.startswith("inword")


PAGE_HINT = re.compile(r"(?i)\bpage\b")
ADDRESS_HINT = re.compile(
    r"(?i)\b(?:"
    r"road|rd\.?|street|st\.?|marg|nagar|link|"
    r"west|east|north|south|floor|office|plot|sector|"
    r"mumbai|malad|delhi|gurgaon|gurugram|maharashtra|haryana|india|"
    r"pincode|pin code|pin|"
    r"bill\s+to|ship\s+to"
    r")\b"
)
PROTECTED_STOP = {
    "page",
    "road",
    "mumbai",
    "malad",
    "maharashtra",
    "delhi",
    "india",
    "pincode",
    "floor",
    "office",
    "bill",
    "ship",
}


def nearby_phrase(words: list[dict], word: dict, px: float = 240, py: float = 18) -> str:
    parts = []
    cx = (word["x0"] + word["x1"]) / 2
    for item in words:
        if abs(item["y0"] - word["y0"]) > py:
            continue
        icx = (item["x0"] + item["x1"]) / 2
        if abs(icx - cx) > px and abs(item["x0"] - word["x0"]) > px:
            continue
        parts.append(item["text"])
    return " ".join(parts)


def is_protected_airtel_text(words: list[dict], word: dict, page) -> bool:
    """Do not mask page numbers or address / pin code text."""
    blob = nearby_phrase(words, word)
    if PAGE_HINT.search(blob):
        return True
    if ADDRESS_HINT.search(blob):
        return True
    digits = re.sub(r"\D", "", word["text"])
    if len(digits) <= 3:
        return True
    return False


def collect_same_line(words: list[dict], start: int, predicate) -> list[dict]:
    if start >= len(words):
        return []
    baseline = words[max(0, start - 1)]["y0"] if start else words[start]["y0"]
    if start < len(words):
        baseline = min(baseline, words[start]["y0"]) if start else words[start]["y0"]
    hits = []
    index = start
    while index < len(words):
        word = words[index]
        if word["y0"] - baseline > 12 and hits:
            break
        if word["y0"] - baseline > 12 and not hits:
            break
        if word["text"].strip() in {":", "-"}:
            index += 1
            continue
        if predicate(word["text"]):
            hits.append(word)
            baseline = min(baseline, word["y0"])
        elif hits:
            break
        index += 1
    return hits


def collect_wrapped(words: list[dict], start: int, predicate, stop_on_label=True) -> list[dict]:
    hits = []
    index = start
    last_y = words[start - 1]["y0"] if start else 0
    while index < len(words):
        word = words[index]
        token = norm_token(word["text"])
        if stop_on_label and token in {
            "payment",
            "special",
            "instructions",
            "unit",
            "qty",
            "quantity",
            *PROTECTED_STOP,
        }:
            break
        if word["y0"] - last_y > 28 and hits:
            break
        if word["text"].strip() in {":", "-"} and not hits:
            index += 1
            continue
        if is_inr_token(word["text"]) and predicate is is_amount_word_token:
            last_y = word["y0"]
            index += 1
            continue
        if predicate(word["text"]):
            hits.append(word)
            last_y = word["y0"]
        elif hits and abs(word["y0"] - last_y) <= 8:
            break
        elif hits:
            break
        index += 1
    return hits


def collect_terms(words: list[dict], start: int) -> list[dict]:
    hits = []
    index = start
    last_y = words[start - 1]["y0"] if start else 0
    while index < len(words):
        word = words[index]
        token = norm_token(word["text"])
        if token in {"special", "instructions", *PROTECTED_STOP}:
            break
        if word["text"].strip() in {":", "-"} and not hits:
            index += 1
            continue
        if word["y0"] - last_y > 22 and hits:
            break
        hits.append(word)
        last_y = word["y0"]
        index += 1
    return hits


def union_word_rect(items: list[dict]) -> fitz.Rect | None:
    if not items:
        return None
    rect = fitz.Rect(items[0]["rect"])
    for item in items[1:]:
        rect |= item["rect"]
    return rect


def add_star_group(hits: list, items: list[dict]) -> None:
    """One 5-star mask for a whole value, not stars on every word."""
    if not items:
        return
    rows = group_word_rows(items, tol=4)
    first_rect = union_word_rect(rows[0])
    if first_rect:
        hits.append((first_rect, GROUP_STARS))
    for row in rows[1:]:
        extra = union_word_rect(row)
        if extra:
            hits.append((extra, ""))


def find_airtel_label_rects(page) -> list[tuple[fitz.Rect, str]]:
    words = page_word_items(page)
    if not words:
        return []
    hits = []
    used = set()

    for start in range(len(words)):
        if start in used:
            continue
        matched = False
        for seq in AIRTEL_WORDS_LABELS:
            end = match_label_at(words, start, seq)
            if end is None:
                continue
            add_star_group(hits, collect_wrapped(words, end, is_amount_word_token))
            used.update(range(start, end))
            matched = True
            break
        if matched:
            continue
        for seq in AIRTEL_MONEY_LABELS:
            end = match_label_at(words, start, seq)
            if end is None:
                continue
            if rest_has_in_words(words, end):
                continue
            add_star_group(hits, collect_same_line(words, end, is_money_token))
            used.update(range(start, end))
            matched = True
            break
        if matched:
            continue
        for seq in AIRTEL_TERMS_LABELS:
            end = match_label_at(words, start, seq)
            if end is None:
                continue
            add_star_group(hits, collect_terms(words, end))
            used.update(range(start, end))
            break

    for row in group_word_rows(words):
        if not any(GST_LABEL.match(word["text"].replace(" ", "")) for word in row):
            continue
        money_cells = [
            word
            for word in row
            if is_money_token(word["text"]) and not is_protected_airtel_text(words, word, page)
        ]
        add_star_group(hits, money_cells)
    return hits


def airtel_header_kind(parts: list[str]) -> str | None:
    phrase = " ".join(parts).lower()
    phrase = re.sub(r"[^a-z0-9\s]+", " ", phrase)
    phrase = " ".join(phrase.split())
    if phrase == "unit price":
        return "unit price"
    if phrase == "line total":
        return "line total"
    return None


def detect_airtel_columns(items: list[dict]):
    """Find Unit Price / Line Total columns from a header row on this page."""
    for row in group_word_rows(items, tol=12):
        joined = " ".join(word["text"] for word in row).lower()
        if "unit" not in joined or "price" not in joined or "line" not in joined or "total" not in joined:
            continue
        columns = []
        header_word_bottoms = []
        index = 0
        while index < len(row):
            one = [row[index]["text"]]
            two = one + [row[index + 1]["text"]] if index + 1 < len(row) else None
            kind = airtel_header_kind(two) if two else None
            span = 2 if kind else 1
            if not kind:
                kind = airtel_header_kind(one)
                span = 1
            if kind:
                columns.append([row[index]["x0"], row[index + span - 1]["x1"], kind])
                header_word_bottoms.extend(
                    row[j]["y1"] for j in range(index, index + span)
                )
                index += span
            else:
                index += 1
        if not columns:
            continue
        sized = []
        for i, (left, right, kind) in enumerate(columns):
            next_left = columns[i + 1][0] if i + 1 < len(columns) else None
            col_right = ((right + next_left) / 2) if next_left is not None else right + 90
            sized.append((left - 8, col_right, kind))
        # Only the Unit Price / Line Total words set the cutoff. Other header
        # cells wrap onto a second line (e.g. "End Date") and would otherwise
        # push the cutoff into the first data row.
        header_bottom = max(header_word_bottoms) if header_word_bottoms else max(
            word["y1"] for word in row
        )
        return sized, header_bottom
    return None


def page_has_airtel_table_signal(items: list[dict]) -> bool:
    """Continuation pages have no header, but do have IGST / Total Line Value rows."""
    for row in group_word_rows(items):
        if any(GST_LABEL.match(word["text"].replace(" ", "")) for word in row):
            return True
        norms = [norm_token(word["text"]) for word in row]
        if norms[:3] == ["total", "line", "value"] or "totallinevalue" in "".join(norms):
            return True
    return False


def find_airtel_column_rects(page, carry_columns=None):
    """Mask Unit Price / Line Total cells. Returns (hits, columns for next pages).

    Airtel continuation sheets (page 2+) do not repeat the table header, so
    columns found on an earlier page are carried forward to those pages.
    """
    items = page_word_items(page)
    if not items:
        return [], carry_columns
    detected = detect_airtel_columns(items)
    if detected:
        sized, header_bottom = detected
        active_columns = sized
    elif carry_columns and page_has_airtel_table_signal(items):
        sized = carry_columns
        active_columns = carry_columns
        header_bottom = float(page.rect.y0)
    else:
        return [], carry_columns

    hits = []
    footer_cut = page.rect.y0 + float(page.rect.height) * 0.86
    found_words = []
    for word in items:
        if word["y0"] < header_bottom - 1:
            continue
        if word["y0"] > footer_cut:
            continue
        if not is_money_token(word["text"]):
            continue
        digits = re.sub(r"\D", "", word["text"])
        if not digits:
            continue
        if is_protected_airtel_text(items, word, page):
            continue
        center_x = (word["x0"] + word["x1"]) / 2
        matched = False
        for left, right, _kind in sized:
            if left <= center_x <= right:
                matched = True
                break
        if not matched:
            # Wide amounts can start before the column centre; accept overlap.
            for left, right, _kind in sized:
                if word["x1"] > left and word["x0"] < right:
                    matched = True
                    break
        if matched:
            found_words.append(word)
    # Group wrapped continuations (e.g. "322548.4800000" / "0004") so the
    # full amount gets one ***** on its first line, not a star per fragment.
    used_ids = set()
    for word in found_words:
        if id(word) in used_ids:
            continue
        group = [word]
        for other in found_words:
            if other is word or id(other) in used_ids:
                continue
            dy = other["y0"] - word["y0"]
            if dy < 4 or dy > 18:
                continue
            if not (other["x0"] < word["x1"] and other["x1"] > word["x0"]):
                continue
            token = (other["text"] or "").replace(" ", "").strip()
            is_fragment = bool(
                re.fullmatch(r"\.\d+", token) or re.fullmatch(r"\d{1,6}", token)
            )
            if is_fragment:
                group.append(other)
                used_ids.add(id(other))
                break
        used_ids.add(id(word))
        add_star_group(hits, group)
    return hits, active_columns


def unique_rects(items: list[tuple[fitz.Rect, str]]) -> list[tuple[fitz.Rect, str]]:
    unique = []
    seen = set()
    for rect, stars in items:
        key = (round(rect.x0, 1), round(rect.y0, 1), round(rect.x1, 1), round(rect.y1, 1))
        if key not in seen:
            seen.add(key)
            unique.append((rect, stars))
    return unique


def find_airtel_rects(page, carry_columns=None):
    column_hits, active_columns = find_airtel_column_rects(page, carry_columns)
    hits = unique_rects(find_airtel_label_rects(page) + column_hits)
    return hits, active_columns


def find_beetal_rects(page) -> list[tuple[fitz.Rect, str]]:
    return unique_rects(find_beetal_label_rects(page) + find_beetal_column_rects(page))


DATE_TOKEN = re.compile(r"^\d{1,2}[./-]\d{1,2}[./-]\d{2,4}$")
WRAPPED_DEC = re.compile(r"^\.\d{1,3}$")
BEETAL_RATE = re.compile(r"^[\d,]+(?:\.\d{1,3})?$")

BEETAL_WORDS_LABELS = (
    ("po", "values", "in", "words"),
    ("po", "value", "in", "words"),
)
BEETAL_TOTAL_VALUE_LABELS = (
    ("total", "value", "with", "gst", "tax"),
    ("total", "value", "withgst", "tax"),
)
BEETAL_TERMS_LABELS = (("payment", "terms"),)


def is_beetal_column_token(text: str, kind: str) -> bool:
    raw = (text or "").replace(" ", "").replace("₹", "").strip()
    token = raw.rstrip(".")
    if kind == "date":
        return bool(DATE_TOKEN.fullmatch(raw) or DATE_TOKEN.fullmatch(token))
    if WRAPPED_DEC.fullmatch(raw):
        return True
    if BEETAL_RATE.fullmatch(token) or BEETAL_RATE.fullmatch(raw):
        return True
    return is_money_token(token) or is_money_token(raw)


def beetal_header_kind(parts: list[str]) -> str | None:
    phrase = " ".join(parts).lower()
    phrase = re.sub(r"[^a-z0-9\s]+", " ", phrase)
    phrase = " ".join(phrase.split())
    if phrase in {"unit price"}:
        return "money"
    if phrase in {"delivery date"}:
        return "date"
    if phrase in {"total amount"}:
        return "money"
    if phrase == "amount":
        return "money"
    if phrase == "rate":
        return "rate"
    return None


def colon_value_rect(word: dict) -> fitz.Rect | None:
    text = word["text"] or ""
    pos = text.find(":")
    if pos < 0:
        return None
    value = text[pos + 1 :].strip()
    if len(value) < 2:
        return None
    width = max(word["x1"] - word["x0"], 1.0)
    x0 = word["x0"] + width * ((pos + 1) / max(len(text), 1))
    x0 = min(word["x1"] - 10, max(word["x0"] + 8, x0))
    return fitz.Rect(x0, word["y0"], word["x1"], word["y1"])


def find_beetal_column_rects(page) -> list[tuple[fitz.Rect, str]]:
    items = page_word_items(page)
    if not items:
        return []
    hits = []
    for row in group_word_rows(items):
        joined = " ".join(word["text"] for word in row).lower()
        if "unit" not in joined or "price" not in joined:
            continue
        if "amount" not in joined and "delivery" not in joined:
            continue
        columns = []
        index = 0
        while index < len(row):
            one = [row[index]["text"]]
            two = one + [row[index + 1]["text"]] if index + 1 < len(row) else None
            kind = beetal_header_kind(two) if two else None
            span = 2 if kind else 1
            if not kind:
                kind = beetal_header_kind(one)
                span = 1
            if kind:
                columns.append([row[index]["x0"], row[index + span - 1]["x1"], kind])
                index += span
            else:
                index += 1
        if not columns:
            continue
        sized = []
        for i, (left, right, kind) in enumerate(columns):
            next_left = columns[i + 1][0] if i + 1 < len(columns) else None
            prev_right = columns[i - 1][1] if i else left
            col_left = (prev_right + left) / 2 if i else left - 14
            col_right = ((right + next_left) / 2) if next_left is not None else right + 100
            sized.append((col_left, col_right, kind))
        header_bottom = max(word["y1"] for word in row)
        footer_cut = page.rect.y0 + float(page.rect.height) * 0.92
        found_words = []
        for word in items:
            if word["y0"] < header_bottom - 1 or word["y0"] > footer_cut:
                continue
            if PAGE_HINT.search(nearby_phrase(items, word, px=80, py=8)):
                continue
            if WRAPPED_DEC.fullmatch((word["text"] or "").strip()):
                continue
            matched = False
            for left, right, kind in sized:
                overlaps = word["x1"] > left and word["x0"] < right
                if overlaps and is_beetal_column_token(word["text"], kind):
                    found_words.append(word)
                    matched = True
                    break
            if matched:
                continue
            if is_beetal_column_token(word["text"], "money"):
                for left, right, kind in sized:
                    if kind in {"money", "rate"} and word["x1"] > left and word["x0"] < right + 40:
                        found_words.append(word)
                        break
        used = set()
        for word in found_words:
            if id(word) in used:
                continue
            group = [word]
            for other in items:
                if other is word or id(other) in used:
                    continue
                if not WRAPPED_DEC.fullmatch((other["text"] or "").strip()):
                    continue
                if other["y0"] < word["y0"] + 1:
                    continue
                if other["y0"] - word["y1"] > 14:
                    continue
                if other["x0"] < word["x1"] + 8 and other["x1"] > word["x0"] - 8:
                    group.append(other)
                    used.add(id(other))
                    break
            used.add(id(word))
            add_star_group(hits, group)
    return hits


def find_beetal_payment_rects(page) -> list[tuple[fitz.Rect, str]]:
    words = page_word_items(page)
    if not words:
        return []
    hits = []
    text = page.get_text("text")
    for match in re.finditer(r"(?i)payment\s*terms\s*:?\s*(.+)", text):
        value = match.group(1).splitlines()[0].strip()
        value = re.split(r"(?i)\b(?:currency|ship\s*method|po\s+values?)\b", value)[0].strip(" :-")
        if len(value) >= 2:
            for snippet in (value, " ".join(value.split()[:4])):
                for rect in page.search_for(snippet):
                    hits.append((rect, GROUP_STARS))
    for word in words:
        if "payment" in word["text"].lower() and "terms" in word["text"].lower() and ":" in word["text"]:
            rect = colon_value_rect(word)
            if rect:
                hits.append((rect, GROUP_STARS))
    for row in group_word_rows(words):
        joined = " ".join(item["text"] for item in row)
        if not re.search(r"(?i)payment\s*terms", joined):
            continue
        colon_x = None
        for item in row:
            if ":" in item["text"]:
                if "payment" in item["text"].lower():
                    rect = colon_value_rect(item)
                    if rect:
                        hits.append((rect, GROUP_STARS))
                colon_x = item["x1"] if item["text"].strip() == ":" else (item["x0"] + item["x1"]) / 2
                if ":" in item["text"]:
                    pos = item["text"].find(":")
                    width = max(item["x1"] - item["x0"], 1)
                    colon_x = item["x0"] + width * ((pos + 1) / max(len(item["text"]), 1))
                break
        value_words = [item for item in row if colon_x is not None and item["x0"] >= colon_x - 2]
        value_words = [
            item
            for item in value_words
            if item["text"].strip() not in {":"} and not re.search(r"(?i)payment\s*terms", item["text"])
        ]
        add_star_group(hits, value_words)
    return hits


def find_beetal_label_rects(page) -> list[tuple[fitz.Rect, str]]:
    words = page_word_items(page)
    if not words:
        return []
    hits = []
    used = set()
    for start in range(len(words)):
        if start in used:
            continue
        matched = False
        for seq in BEETAL_WORDS_LABELS:
            end = match_label_at(words, start, seq)
            if end is None:
                continue
            add_star_group(hits, collect_wrapped(words, end, is_amount_word_token))
            used.update(range(start, end))
            matched = True
            break
        if matched:
            continue
        for seq in BEETAL_TOTAL_VALUE_LABELS:
            end = match_label_at(words, start, seq)
            if end is None:
                continue
            add_star_group(hits, collect_same_line(words, end, lambda t: is_beetal_column_token(t, "money")))
            used.update(range(start, end))
            matched = True
            break
        if matched:
            continue
        for seq in BEETAL_TERMS_LABELS:
            end = match_label_at(words, start, seq)
            if end is None:
                continue
            add_star_group(hits, collect_terms(words, end))
            used.update(range(start, end))
            break
    hits.extend(find_beetal_payment_rects(page))
    return hits


def paint_bounds(selection: fitz.Rect, word_rects: list[fitz.Rect]) -> fitz.Rect:
    """Keep stars on the selected line only, not the line below."""
    if not word_rects:
        trim = min(3.0, selection.height * 0.25)
        return fitz.Rect(selection.x0, selection.y0, selection.x1, max(selection.y0 + 1, selection.y1 - trim))
    union = word_rects[0]
    for rect in word_rects[1:]:
        union |= rect
    return fitz.Rect(selection.x0, union.y0, selection.x1, union.y1)


def mask_pdf(
    source_path: Path,
    dest_path: Path,
    selected_types: list[str],
    extra_terms: list[str],
    auto: bool,
    vendor: str = "airtel",
) -> dict:
    doc = fitz.open(source_path)
    total_redactions = 0
    pages_touched = 0
    found_terms: list[str] = []
    vendor = (vendor or "airtel").lower().strip()

    airtel_carry_columns = None
    for page in doc:
        items = []
        if auto and vendor == "airtel":
            items, airtel_carry_columns = find_airtel_rects(page, airtel_carry_columns)
        elif auto and vendor == "beetal":
            items = find_beetal_rects(page)
        else:
            text = page.get_text("text")
            matches = find_matches(text, selected_types, extra_terms, auto)
            found_terms.extend(original for original, _replacement in matches)
            for original, replacement in matches:
                for rect in page.search_for(original):
                    items.append((rect, replacement))
            if auto:
                items.extend(find_table_amount_rects(page))
        if items:
            total_redactions += redact_with_stars(page, items)
            pages_touched += 1

    doc.save(dest_path, garbage=4, deflate=True)
    page_count = doc.page_count
    doc.close()

    note = None

    return {
        "pages": page_count,
        "pages_masked": pages_touched,
        "redactions": total_redactions,
        "mode": "auto" if auto else "manual",
        "vendor": vendor,
        "note": note,
        "detected": unique_terms(found_terms)[:20],
    }


def is_job_id(job_id: str) -> bool:
    return bool(JOB_ID_RE.fullmatch(job_id or ""))


def save_upload(uploaded) -> tuple[str, Path]:
    job_id = uuid.uuid4().hex
    source_path = UPLOAD_DIR / f"{job_id}.pdf"
    uploaded.save(source_path)
    return job_id, source_path


def render_page_images(source_path: Path, job_id: str) -> list[dict]:
    job_dir = PAGES_DIR / job_id
    job_dir.mkdir(parents=True, exist_ok=True)
    doc = fitz.open(source_path)
    pages = []
    matrix = fitz.Matrix(PAGE_ZOOM, PAGE_ZOOM)
    for index, page in enumerate(doc):
        pix = page.get_pixmap(matrix=matrix, clip=page.rect, alpha=False)
        pix.save(job_dir / f"{index}.png")
        pages.append({"index": index, "url": f"/page/{job_id}/{index}"})
    doc.close()
    return pages


def mask_regions(source_path: Path, dest_path: Path, regions: list[dict]) -> dict:
    grouped: dict[int, list[dict]] = {}
    for region in regions:
        grouped.setdefault(int(region["page"]), []).append(region)

    doc = fitz.open(source_path)
    total_redactions = 0
    pages_touched = 0

    for page_index, page in enumerate(doc):
        page_regions = grouped.get(page_index, [])
        if not page_regions:
            continue
        items = []
        paint_rects = []
        for region in page_regions:
            rect = display_rect(page, region)
            if rect.width < 1 or rect.height < 1:
                continue
            word_rects = words_in_selection(page, rect)
            items.extend((word_rect, None) for word_rect in word_rects)
            paint_rects.append(paint_bounds(rect, word_rects))
        if items:
            total_redactions += redact_with_stars(page, items, fill_stars=False)
            pages_touched += 1
            for paint in paint_rects:
                if paint.width > 1 and paint.height > 1:
                    write_stars(page, paint, single_line=True)
        elif paint_rects:
            for paint in paint_rects:
                if paint.width > 1 and paint.height > 1:
                    write_stars(page, paint, single_line=True)
                    total_redactions += 1
            pages_touched += 1

    doc.save(dest_path, garbage=4, deflate=True)
    page_count = doc.page_count
    doc.close()
    return {
        "pages": page_count,
        "pages_masked": pages_touched,
        "redactions": total_redactions,
        "mode": "select",
    }


@app.route("/")
def index():
    return render_template("index.html")


def safe_download_name(original: str) -> str:
    stem = Path(original or "document").stem
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", stem).strip("._") or "document"
    return f"{cleaned[:80]}_masked.pdf"


@app.post("/api/mask")
def api_mask():
    uploaded = request.files.get("file") or (request.files.getlist("files") or [None])[0]
    if not uploaded or not uploaded.filename:
        return jsonify({"error": "Please upload a PDF file."}), 400

    filename = uploaded.filename.lower()
    if not filename.endswith(".pdf"):
        return jsonify({"error": "Only PDF files are supported."}), 400

    auto = request.form.get("mode", "auto") == "auto"
    vendor = request.form.get("vendor", "airtel")
    selected = request.form.getlist("types") or DEFAULT_TYPES
    extra_raw = request.form.get("keywords", "")
    extra_terms = [term.strip() for term in extra_raw.split(",") if term.strip()]

    job_id, source_path = save_upload(uploaded)
    dest_path = OUTPUT_DIR / f"{job_id}.pdf"

    try:
        stats = mask_pdf(source_path, dest_path, selected, extra_terms, auto, vendor)
    except Exception as exc:
        return jsonify({"error": f"Could not process this PDF: {exc}"}), 500

    return jsonify(
        {
            "job_id": job_id,
            "original_name": uploaded.filename,
            "preview_url": f"/preview/{job_id}",
            "download_url": f"/download/{job_id}",
            **stats,
        }
    )


@app.post("/api/mask-batch")
def api_mask_batch():
    uploads = request.files.getlist("files") or request.files.getlist("file")
    uploads = [item for item in uploads if item and item.filename]
    if not uploads:
        return jsonify({"error": "Please upload at least one PDF file."}), 400
    if len(uploads) > MAX_FILES:
        return jsonify({"error": f"Upload up to {MAX_FILES} PDFs at once."}), 400

    vendor = (request.form.get("vendor") or "airtel").lower().strip()
    if vendor not in {"airtel", "beetal"}:
        return jsonify({"error": "Choose Airtel or Beetal."}), 400

    results = []
    batch_id = uuid.uuid4().hex
    zip_path = OUTPUT_DIR / f"{batch_id}.zip"

    try:
        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            used_names = set()
            for uploaded in uploads:
                if not uploaded.filename.lower().endswith(".pdf"):
                    return jsonify({"error": f"{uploaded.filename} is not a PDF."}), 400
                job_id, source_path = save_upload(uploaded)
                dest_path = OUTPUT_DIR / f"{job_id}.pdf"
                stats = mask_pdf(source_path, dest_path, DEFAULT_TYPES, [], True, vendor)
                download_name = safe_download_name(uploaded.filename)
                if download_name in used_names:
                    download_name = f"{job_id[:8]}_{download_name}"
                used_names.add(download_name)
                archive.write(dest_path, arcname=download_name)
                results.append(
                    {
                        "job_id": job_id,
                        "original_name": uploaded.filename,
                        "preview_url": f"/preview/{job_id}",
                        "download_url": f"/download/{job_id}",
                        **stats,
                    }
                )
    except Exception as exc:
        return jsonify({"error": f"Could not process these PDFs: {exc}"}), 500

    note = results[0].get("note") if results else None
    total_redactions = sum(item.get("redactions", 0) for item in results)
    return jsonify(
        {
            "batch_id": batch_id,
            "vendor": vendor,
            "mode": "auto",
            "note": note,
            "redactions": total_redactions,
            "zip_url": f"/download-zip/{batch_id}",
            "files": results,
        }
    )


@app.post("/api/upload")
def api_upload():
    uploaded = request.files.get("file")
    if not uploaded or not uploaded.filename:
        return jsonify({"error": "Please upload a PDF file."}), 400
    if not uploaded.filename.lower().endswith(".pdf"):
        return jsonify({"error": "Only PDF files are supported."}), 400

    job_id, source_path = save_upload(uploaded)
    try:
        pages = render_page_images(source_path, job_id)
    except Exception as exc:
        return jsonify({"error": f"Could not open this PDF: {exc}"}), 500

    return jsonify({"job_id": job_id, "original_name": uploaded.filename, "pages": pages})


@app.post("/api/mask-regions")
def api_mask_regions():
    payload = request.get_json(silent=True) or {}
    job_id = payload.get("job_id", "")
    regions = payload.get("regions") or []
    if not is_job_id(job_id):
        return jsonify({"error": "Invalid file id."}), 400

    source_path = UPLOAD_DIR / f"{job_id}.pdf"
    dest_path = OUTPUT_DIR / f"{job_id}.pdf"
    if not source_path.exists():
        return jsonify({"error": "Uploaded PDF not found. Please upload again."}), 404

    try:
        stats = mask_regions(source_path, dest_path, regions)
    except Exception as exc:
        return jsonify({"error": f"Could not mask this PDF: {exc}"}), 500

    return jsonify(
        {
            "job_id": job_id,
            "preview_url": f"/preview/{job_id}",
            "download_url": f"/download/{job_id}",
            **stats,
        }
    )


@app.get("/page/<job_id>/<int:page_index>")
def serve_page(job_id: str, page_index: int):
    if not is_job_id(job_id) or page_index < 0:
        return jsonify({"error": "Invalid page."}), 400
    path = PAGES_DIR / job_id / f"{page_index}.png"
    if not path.exists():
        return jsonify({"error": "Page not found."}), 404
    return send_file(path, mimetype="image/png")


@app.get("/preview/<job_id>")
@app.get("/download/<job_id>")
def serve_masked(job_id: str):
    if not is_job_id(job_id):
        return jsonify({"error": "Invalid file id."}), 400

    dest_path = OUTPUT_DIR / f"{job_id}.pdf"
    if not dest_path.exists():
        return jsonify({"error": "Masked PDF not found. Please upload again."}), 404

    as_attachment = request.path.startswith("/download/")
    original = request.args.get("name") or "masked.pdf"
    return send_file(
        dest_path,
        mimetype="application/pdf",
        as_attachment=as_attachment,
        download_name=original if original.lower().endswith(".pdf") else "masked.pdf",
    )


@app.get("/download-zip/<batch_id>")
def serve_zip(batch_id: str):
    if not is_job_id(batch_id):
        return jsonify({"error": "Invalid file id."}), 400
    zip_path = OUTPUT_DIR / f"{batch_id}.zip"
    if not zip_path.exists():
        return jsonify({"error": "Zip file not found. Please mask again."}), 404
    return send_file(
        zip_path,
        mimetype="application/zip",
        as_attachment=True,
        download_name="masked_pdfs.zip",
    )


def is_pdf_name(name: str) -> bool:
    return (name or "").lower().endswith(".pdf")


def is_word_name(name: str) -> bool:
    lower = (name or "").lower()
    return lower.endswith(".docx") or lower.endswith(".doc")


def convert_stem(original: str) -> str:
    stem = Path(original or "document").stem
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", stem).strip("._") or "document"
    return cleaned[:80]


def find_soffice() -> str | None:
    which = shutil.which("soffice") or shutil.which("soffice.exe")
    if which:
        return which
    candidates = [
        Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
        Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
    ]
    for path in candidates:
        if path.exists():
            return str(path)
    return None


AMOUNT_CELL_RE = re.compile(
    r"^[\d,]+(?:\.\d{1,3})?$|^[\d,]+\.$|^\.\d{1,3}$|^\d{1,2}[./-]\d{1,2}[./-]\d{2,4}$"
)
AMOUNT_HEADER_RE = re.compile(
    r"(?i)\b(unit\s*price|amount|total(?:\s*amount)?|rate|price|value)\b"
)


def _cell_plain_text(cell) -> str:
    parts = []
    for paragraph in cell.paragraphs:
        parts.append("".join(run.text or "" for run in paragraph.runs) or (paragraph.text or ""))
    return "".join(parts).strip()


def _looks_like_amount_fragment(text: str) -> bool:
    compact = re.sub(r"\s+", "", text or "")
    if not compact:
        return False
    if AMOUNT_CELL_RE.fullmatch(compact):
        return True
    # Broken cells often store one digit / comma / decimal per paragraph.
    joined = "".join(compact.split())
    return bool(re.fullmatch(r"[\d,.]+", joined) and re.search(r"\d", joined) and len(joined) <= 24)


def _set_cell_nowrap(cell) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    no_wrap = tc_pr.find(qn("w:noWrap"))
    if no_wrap is None:
        no_wrap = OxmlElement("w:noWrap")
        tc_pr.append(no_wrap)


def _rewrite_cell_single_line(cell, text: str) -> None:
    """Replace broken multi-line cell content with one readable amount line."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    first = cell.paragraphs[0]
    for paragraph in list(cell.paragraphs[1:]):
        element = paragraph._element
        element.getparent().remove(element)

    for run in list(first.runs):
        run._element.getparent().remove(run._element)
    first.text = ""
    run = first.add_run(text)
    run.font.size = Pt(8)
    first.paragraph_format.space_before = Pt(0)
    first.paragraph_format.space_after = Pt(0)
    first.paragraph_format.line_spacing = 1.0

    p_pr = first._p.get_or_add_pPr()
    word_wrap = p_pr.find(qn("w:wordWrap"))
    if word_wrap is None:
        word_wrap = OxmlElement("w:wordWrap")
        p_pr.append(word_wrap)
    word_wrap.set(qn("w:val"), "0")
    _set_cell_nowrap(cell)

    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    needed = max(1400, min(3200, 150 * max(len(text), 6)))  # twips
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    try:
        current = int(tc_w.get(qn("w:w")) or "0")
    except Exception:
        current = 0
    if current < needed:
        tc_w.set(qn("w:w"), str(needed))
        tc_w.set(qn("w:type"), "dxa")


def fix_docx_collapsed_amounts(docx_path: Path) -> None:
    """Repair amount/price cells that pdf2docx wrapped into vertical digit stacks."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document(str(docx_path))
    for table in document.tables:
        table.autofit = True
        table.allow_autofit = True
        tbl = table._tbl
        tbl_pr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
        if tbl.tblPr is None:
            tbl.insert(0, tbl_pr)
        layout = tbl_pr.find(qn("w:tblLayout"))
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            tbl_pr.append(layout)
        layout.set(qn("w:type"), "autofit")

        header_amount_cols: set[int] = set()
        if table.rows:
            for idx, cell in enumerate(table.rows[0].cells):
                if AMOUNT_HEADER_RE.search(cell.text or ""):
                    header_amount_cols.add(idx)

        for row in table.rows:
            for col_idx, cell in enumerate(row.cells):
                raw_paragraphs = [(p.text or "").strip() for p in cell.paragraphs]
                compact = re.sub(r"\s+", "", "".join(raw_paragraphs))
                non_empty = [part for part in raw_paragraphs if part]
                many_tiny_lines = len(non_empty) >= 2 and all(len(part) <= 4 for part in non_empty)
                is_amount_col = col_idx in header_amount_cols
                if not (
                    is_amount_col
                    or _looks_like_amount_fragment(compact)
                    or (many_tiny_lines and re.fullmatch(r"[\d,.]+", compact or ""))
                ):
                    continue
                if not compact or not re.search(r"\d", compact):
                    continue
                _rewrite_cell_single_line(cell, compact)

    document.save(str(docx_path))


def pdf_to_word(source: Path, dest: Path) -> None:
    """Convert PDF to editable Word, then fix crushed amount columns."""
    converter = Converter(str(source))
    try:
        converter.convert(
            str(dest),
            parse_lattice_table=True,
            parse_stream_table=True,
            extract_stream_table=False,
            min_section_height=8.0,
            page_margin_factor_top=1.0,
            page_margin_factor_bottom=1.0,
            connected_border_tolerance=1.5,
            max_border_width=10.0,
            min_border_clearance=1.0,
            clip_image_res_ratio=5.0,
            float_image_ignorable_gap=2.0,
            line_separate_threshold=2.5,
            line_break_width_ratio=0.25,
            line_break_free_space_ratio=0.05,
            new_paragraph_free_space_ratio=0.95,
            lines_left_aligned_threshold=1.5,
            lines_right_aligned_threshold=1.5,
            delete_end_line_hyphen=False,
            ignore_page_error=True,
        )
    finally:
        converter.close()
    if not dest.exists() or dest.stat().st_size == 0:
        raise RuntimeError("PDF → Word conversion produced an empty file.")
    try:
        fix_docx_collapsed_amounts(dest)
    except Exception:
        # Conversion succeeded even if cleanup fails; keep the editable Word file.
        pass


# ---------------------------------------------------------------------------
# Built-in Word (.docx) -> PDF conversion. No Microsoft Word / LibreOffice
# installation needed: the document is rendered with PyMuPDF.
# ---------------------------------------------------------------------------

@lru_cache(maxsize=8192)
def _text_width(token: str, fontname: str, fontsize: float) -> float:
    try:
        return fitz.get_text_length(token, fontname=fontname, fontsize=fontsize)
    except Exception:
        return len(token) * fontsize * 0.5


_DOCX_FONT_MAP = {
    "courier": ("cour", "cobo", "coit", "cobi"),
    "consolas": ("cour", "cobo", "coit", "cobi"),
    "times": ("tiro", "tibo", "tiit", "tibi"),
    "georgia": ("tiro", "tibo", "tiit", "tibi"),
    "garamond": ("tiro", "tibo", "tiit", "tibi"),
}


def _docx_font(run, base_size: float) -> tuple[str, float]:
    name = ""
    size = base_size
    try:
        if run.font.size:
            size = float(run.font.size.pt)
        if run.font.name:
            name = run.font.name.lower()
        elif run.style is not None and run.style.font.name:
            name = run.style.font.name.lower()
    except Exception:
        pass
    family = ("helv", "hebo", "heit", "hebi")
    for key, fonts in _DOCX_FONT_MAP.items():
        if key in name:
            family = fonts
            break
    bold = bool(run.bold)
    italic = bool(run.italic)
    if bold and italic:
        return family[3], size
    if bold:
        return family[1], size
    if italic:
        return family[2], size
    return family[0], size


def _docx_color(run) -> tuple[float, float, float]:
    try:
        rgb = run.font.color.rgb
        if rgb is not None:
            return (rgb[0] / 255, rgb[1] / 255, rgb[2] / 255)
    except Exception:
        pass
    return (0, 0, 0)


class _DocxPdfRenderer:
    def __init__(self, document):
        self.document = document
        section = document.sections[0]
        self.page_w = float(section.page_width.pt)
        self.page_h = float(section.page_height.pt)
        self.margin_l = float(section.left_margin.pt) if section.left_margin else 36.0
        self.margin_r = float(section.right_margin.pt) if section.right_margin else 36.0
        self.margin_t = float(section.top_margin.pt) if section.top_margin else 36.0
        self.margin_b = float(section.bottom_margin.pt) if section.bottom_margin else 36.0
        self.content_w = max(100.0, self.page_w - self.margin_l - self.margin_r)
        self.base_size = 10.5
        try:
            normal = document.styles["Normal"]
            if normal.font.size:
                self.base_size = float(normal.font.size.pt)
        except Exception:
            pass
        self.pdf = fitz.open()
        self.page = None
        self.y = self.margin_t
        self.new_page()

    def new_page(self):
        self.page = self.pdf.new_page(width=self.page_w, height=self.page_h)
        self.y = self.margin_t

    def ensure_space(self, height: float):
        if self.y + height > self.page_h - self.margin_b and self.y > self.margin_t:
            self.new_page()

    def wrap_runs(self, runs, max_width: float):
        """Split styled runs into lines of (text, fontname, size, color)."""
        words = []
        for run in runs:
            text = run.text or ""
            if not text:
                continue
            fontname, size = _docx_font(run, self.base_size)
            color = _docx_color(run)
            for token in re.split(r"(\s+)", text):
                if token:
                    words.append((token, fontname, size, color))
        lines = []
        current = []
        current_w = 0.0
        for token, fontname, size, color in words:
            try:
                width = _text_width(token, fontname, size)
            except Exception:
                width = len(token) * size * 0.5
            if token.strip() == "":
                if current:
                    current.append((token, fontname, size, color))
                    current_w += width
                continue
            if current and current_w + width > max_width:
                lines.append(current)
                current = [(token, fontname, size, color)]
                current_w = width
            else:
                current.append((token, fontname, size, color))
                current_w += width
        if current:
            lines.append(current)
        return lines

    def draw_lines(self, lines, x0: float, max_width: float, alignment) -> float:
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        height = 0.0
        for line in lines:
            line_h = max((size for _t, _f, size, _c in line), default=self.base_size) * 1.28
            line_w = 0.0
            for token, fontname, size, _c in line:
                try:
                    line_w += _text_width(token, fontname, size)
                except Exception:
                    line_w += len(token) * size * 0.5
            if alignment == WD_ALIGN_PARAGRAPH.CENTER:
                start_x = x0 + max(0.0, (max_width - line_w) / 2)
            elif alignment == WD_ALIGN_PARAGRAPH.RIGHT:
                start_x = x0 + max(0.0, max_width - line_w)
            else:
                start_x = x0
            x = start_x
            baseline = self.y + line_h * 0.82
            for token, fontname, size, color in line:
                if token.strip():
                    try:
                        self.page.insert_text(
                            (x, baseline), token, fontname=fontname,
                            fontsize=size, color=color,
                        )
                    except Exception:
                        self.page.insert_text(
                            (x, baseline), token, fontname="helv",
                            fontsize=size, color=color,
                        )
                try:
                    x += _text_width(token, fontname, size)
                except Exception:
                    x += len(token) * size * 0.5
            self.y += line_h
            height += line_h
        return height

    def render_paragraph(self, paragraph):
        text = paragraph.text or ""
        # Embedded images in this paragraph.
        self.render_paragraph_images(paragraph)
        if not text.strip():
            self.y += self.base_size * 0.6
            return
        try:
            space_before = paragraph.paragraph_format.space_before
            space_after = paragraph.paragraph_format.space_after
            if space_before:
                self.y += float(space_before.pt)
        except Exception:
            space_after = None
        try:
            style_name = paragraph.style.name if paragraph.style else ""
        except Exception:
            style_name = ""
        heading_sizes = {"Heading 1": 16.0, "Heading 2": 13.5, "Heading 3": 11.5, "Title": 22.0}
        if style_name in heading_sizes:
            from docx.shared import Pt

            for run in paragraph.runs:
                if run.font.size is None:
                    run.font.size = Pt(heading_sizes[style_name])
                if run.bold is None:
                    run.bold = True
        runs = list(paragraph.runs) or [paragraph]
        if not paragraph.runs:
            # Fallback shim object with .text
            class _R:
                text = paragraph.text
                bold = italic = None
                font = type("F", (), {"size": None, "name": None, "color": type("C", (), {"rgb": None})()})()
                style = None
            runs = [_R()]
        lines = self.wrap_runs(runs, self.content_w)
        line_h_est = max(
            (size for line in lines for _t, _f, size, _c in line),
            default=self.base_size,
        ) * 1.28
        self.ensure_space(min(line_h_est, 40.0))
        self.draw_lines(lines, self.margin_l, self.content_w, paragraph.alignment)
        try:
            if space_after:
                self.y += float(space_after.pt)
            else:
                self.y += 2.0
        except Exception:
            self.y += 2.0

    def render_paragraph_images(self, paragraph):
        from docx.oxml.ns import qn

        for blip in paragraph._p.findall(f".//{qn('a:blip')}"):
            embed = blip.get(qn("r:embed"))
            if not embed:
                continue
            try:
                part = self.document.part.related_parts[embed]
                image_bytes = part.blob
            except Exception:
                continue
            # Image display size if specified.
            width = self.content_w
            extents = paragraph._p.findall(f".//{qn('wp:extent')}")
            if extents:
                try:
                    cx = int(extents[0].get("cx") or 0)
                    if cx:
                        width = min(self.content_w, cx / 12700.0)
                except Exception:
                    pass
            stream = fitz.open("png", image_bytes) if image_bytes[:4] == b"\x89PNG" else None
            try:
                if stream is not None:
                    src_page = stream[0]
                    ratio = src_page.rect.height / max(src_page.rect.width, 1)
                else:
                    pix_doc = fitz.open(stream=image_bytes)
                    src_page = pix_doc[0]
                    ratio = src_page.rect.height / max(src_page.rect.width, 1)
            except Exception:
                ratio = 0.75
            height = width * ratio
            if height > self.page_h - self.margin_t - self.margin_b:
                height = self.page_h - self.margin_t - self.margin_b
                width = height / max(ratio, 0.01)
            self.ensure_space(height)
            try:
                self.page.insert_image(
                    fitz.Rect(self.margin_l, self.y, self.margin_l + width, self.y + height),
                    stream=image_bytes,
                )
            except Exception:
                continue
            self.y += height + 4

    def column_widths(self, table) -> list[float]:
        widths = []
        for column in table.columns:
            try:
                widths.append(float(column.width.pt) if column.width else 0.0)
            except Exception:
                widths.append(0.0)
        if not widths:
            return [self.content_w]
        total = sum(widths)
        if total <= 0:
            return [self.content_w / len(widths)] * len(widths)
        # Normalise to the available page width.
        return [w * self.content_w / total for w in widths]

    def cell_shading(self, cell):
        from docx.oxml.ns import qn

        try:
            tc_pr = cell._tc.get_or_add_tcPr()
            shd = tc_pr.find(qn("w:shd"))
            if shd is not None:
                fill = shd.get(qn("w:fill"))
                if fill and fill.lower() not in {"auto", "none"}:
                    r = int(fill[0:2], 16) / 255
                    g = int(fill[2:4], 16) / 255
                    b = int(fill[4:6], 16) / 255
                    return (r, g, b)
        except Exception:
            pass
        return None

    def render_table(self, table):
        if not table.rows:
            return
        widths = self.column_widths(table)
        col_x = [self.margin_l]
        for w in widths:
            col_x.append(col_x[-1] + w)

        for row in table.rows:
            cells_data = []
            row_h = 8.0
            for idx, cell in enumerate(row.cells):
                max_w = max(30.0, widths[idx] - 7.0)
                cell_lines = []
                for paragraph in cell.paragraphs:
                    if paragraph.runs:
                        cell_lines.extend(self.wrap_runs(list(paragraph.runs), max_w))
                    elif paragraph.text:
                        class _R2:
                            text = paragraph.text
                            bold = italic = None
                            font = type("F", (), {"size": None, "name": None, "color": type("C", (), {"rgb": None})()})()
                            style = None
                        cell_lines.extend(self.wrap_runs([_R2()], max_w))
                cell_h = 5.0
                for line in cell_lines:
                    cell_h += max((s for _t, _f, s, _c in line), default=9.0) * 1.28
                cells_data.append(cell_lines)
                row_h = max(row_h, cell_h)
            self.ensure_space(row_h)
            top = self.y
            for idx, cell in enumerate(row.cells):
                rect = fitz.Rect(col_x[idx], top, col_x[idx + 1], top + row_h)
                shading = self.cell_shading(cell)
                if shading:
                    self.page.draw_rect(rect, color=(0.75, 0.78, 0.82), fill=shading, width=0.4)
                else:
                    self.page.draw_rect(rect, color=(0.75, 0.78, 0.82), width=0.4)
                saved_y = self.y
                self.y = top + 2.5
                # Cell text, left aligned inside the cell.
                for line in cells_data[idx]:
                    line_h = max((s for _t, _f, s, _c in line), default=9.0) * 1.28
                    x = col_x[idx] + 3.5
                    baseline = self.y + line_h * 0.82
                    for token, fontname, size, color in line:
                        if token.strip():
                            try:
                                self.page.insert_text(
                                    (x, baseline), token, fontname=fontname,
                                    fontsize=min(size, 10.0), color=color,
                                )
                            except Exception:
                                self.page.insert_text(
                                    (x, baseline), token, fontname="helv",
                                    fontsize=min(size, 10.0), color=color,
                                )
                        x += _text_width(token, fontname, min(size, 10.0))
                    self.y += line_h
                self.y = saved_y
            self.y = top + row_h
        self.y += 4.0

    def render(self, dest: Path):
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        body = self.document.element.body
        for child in body.iterchildren():
            if child.tag == qn("w:p"):
                self.render_paragraph(Paragraph(child, self.document))
            elif child.tag == qn("w:tbl"):
                self.render_table(Table(child, self.document))
        # Headings style sizes: python-docx runs already carry explicit sizes
        # in most generated files; keep as-is.
        self.pdf.save(str(dest), garbage=4, deflate=True)
        self.pdf.close()


def docx_to_pdf_builtin(source: Path, dest: Path) -> None:
    """Render .docx to PDF in pure Python (no Word / LibreOffice needed)."""
    from docx import Document

    document = Document(str(source))
    renderer = _DocxPdfRenderer(document)
    renderer.render(dest)


def word_to_pdf(source: Path, dest: Path) -> None:
    if source.suffix.lower() == ".docx":
        try:
            docx_to_pdf_builtin(source, dest)
            if dest.exists() and dest.stat().st_size > 0:
                return
        except Exception:
            pass

    try:
        from docx2pdf import convert as docx_convert

        docx_convert(str(source), str(dest))
        if dest.exists():
            return
    except Exception:
        pass

    soffice = find_soffice()
    if not soffice:
        raise RuntimeError(
            "Could not convert this Word file. Built-in conversion supports .docx files; "
            "for old .doc files Microsoft Word or LibreOffice is still needed."
        )

    out_dir = dest.parent
    result = subprocess.run(
        [
            soffice,
            "--headless",
            "--nologo",
            "--nofirststartwizard",
            "--convert-to",
            "pdf",
            "--outdir",
            str(out_dir),
            str(source),
        ],
        capture_output=True,
        text=True,
        timeout=180,
        check=False,
    )
    produced = out_dir / f"{source.stem}.pdf"
    if result.returncode != 0 or not produced.exists():
        detail = (result.stderr or result.stdout or "").strip()
        raise RuntimeError(detail or "LibreOffice could not convert this Word file.")
    if produced != dest:
        produced.replace(dest)


@app.post("/api/convert")
def api_convert():
    direction = (request.form.get("direction") or "").strip().lower()
    if direction not in {"pdf-to-word", "word-to-pdf"}:
        return jsonify({"error": "Choose PDF → Word or Word → PDF."}), 400

    uploads = request.files.getlist("files") or request.files.getlist("file")
    uploads = [item for item in uploads if item and item.filename]
    if not uploads:
        return jsonify({"error": "Please upload at least one file."}), 400
    if len(uploads) > MAX_FILES:
        return jsonify({"error": f"Upload up to {MAX_FILES} files at once."}), 400

    batch_id = uuid.uuid4().hex
    results = []
    zip_path = CONVERT_OUTPUT_DIR / f"{batch_id}.zip"

    try:
        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            used_names = set()
            for uploaded in uploads:
                original = uploaded.filename or "document"
                if direction == "pdf-to-word" and not is_pdf_name(original):
                    return jsonify({"error": f"{original} is not a PDF."}), 400
                if direction == "word-to-pdf" and not is_word_name(original):
                    return jsonify({"error": f"{original} is not a Word file."}), 400

                job_id = uuid.uuid4().hex
                suffix = Path(original).suffix.lower() or (
                    ".pdf" if direction == "pdf-to-word" else ".docx"
                )
                source_path = CONVERT_UPLOAD_DIR / f"{job_id}{suffix}"
                uploaded.save(source_path)

                if direction == "pdf-to-word":
                    output_name = f"{convert_stem(original)}.docx"
                    dest_path = CONVERT_OUTPUT_DIR / f"{job_id}.docx"
                    pdf_to_word(source_path, dest_path)
                    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                else:
                    output_name = f"{convert_stem(original)}.pdf"
                    dest_path = CONVERT_OUTPUT_DIR / f"{job_id}.pdf"
                    word_to_pdf(source_path, dest_path)
                    mime = "application/pdf"

                if output_name in used_names:
                    output_name = f"{job_id[:8]}_{output_name}"
                used_names.add(output_name)
                archive.write(dest_path, arcname=output_name)
                results.append(
                    {
                        "job_id": job_id,
                        "original_name": original,
                        "output_name": output_name,
                        "download_url": f"/download-converted/{job_id}?name={secure_filename(output_name)}",
                        "mime": mime,
                    }
                )
    except Exception as exc:
        return jsonify({"error": f"Could not convert these files: {exc}"}), 500

    payload = {
        "batch_id": batch_id,
        "direction": direction,
        "files": results,
    }
    if len(results) > 1:
        payload["zip_url"] = f"/download-converted-zip/{batch_id}"
    return jsonify(payload)


@app.get("/download-converted/<job_id>")
def serve_converted(job_id: str):
    if not is_job_id(job_id):
        return jsonify({"error": "Invalid file id."}), 400

    docx_path = CONVERT_OUTPUT_DIR / f"{job_id}.docx"
    pdf_path = CONVERT_OUTPUT_DIR / f"{job_id}.pdf"
    if docx_path.exists():
        dest_path = docx_path
        default_name = "converted.docx"
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif pdf_path.exists():
        dest_path = pdf_path
        default_name = "converted.pdf"
        mime = "application/pdf"
    else:
        return jsonify({"error": "Converted file not found. Please convert again."}), 404

    name = request.args.get("name") or default_name
    return send_file(
        dest_path,
        mimetype=mime,
        as_attachment=True,
        download_name=name,
    )


@app.get("/download-converted-zip/<batch_id>")
def serve_converted_zip(batch_id: str):
    if not is_job_id(batch_id):
        return jsonify({"error": "Invalid file id."}), 400
    zip_path = CONVERT_OUTPUT_DIR / f"{batch_id}.zip"
    if not zip_path.exists():
        return jsonify({"error": "Zip file not found. Please convert again."}), 404
    return send_file(
        zip_path,
        mimetype="application/zip",
        as_attachment=True,
        download_name="converted.zip",
    )


if __name__ == "__main__":
    app.run(debug=True, port=5000)
